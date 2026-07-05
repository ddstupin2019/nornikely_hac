import os
import csv
import pandas as pd
from docx import Document as DocxDocument
from PIL import Image
import fitz  # PyMuPDF
import pdfplumber
import pytesseract
from my_solution_v1.core.logger import get_logger

logger = get_logger(__name__)

# Убедимся, что путь к Tesseract установлен
# Если Tesseract установлен в другую директорию, поправьте этот путь
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def parse_pdf(file_path: str) -> str:
    text = ""
    try:
        # Проходим по страницам через pdfplumber для текста и таблиц
        with pdfplumber.open(file_path) as pdf:
            # Если потребуется OCR, откроем параллельно через fitz
            doc_fitz = None
            
            for i, page in enumerate(pdf.pages, 1):
                page_text = ""
                
                # 1. Извлекаем обычный текст
                extracted_text = page.extract_text()
                if extracted_text:
                    page_text += extracted_text + "\n"
                
                # 2. Извлекаем таблицы
                tables = page.extract_tables()
                for table in tables:
                    if table:
                        # Форматируем таблицу как Markdown
                        md_table = ""
                        for row_idx, row in enumerate(table):
                            cleaned_row = [str(cell).replace('\n', ' ') if cell else "" for cell in row]
                            md_table += "| " + " | ".join(cleaned_row) + " |\n"
                            if row_idx == 0:
                                md_table += "|" + "|".join(["---"] * len(row)) + "|\n"
                        page_text += "\n" + md_table + "\n"
                
                # 3. Если текста почти нет (скан), применяем OCR
                if len(page_text.strip()) < 50:
                    logger.info(f"Page {i} in {os.path.basename(file_path)} seems to be a scan. Running OCR...")
                    if not doc_fitz:
                        doc_fitz = fitz.open(file_path)
                    
                    fitz_page = doc_fitz[i - 1]
                    pix = fitz_page.get_pixmap(dpi=150)
                    
                    # Конвертируем pixmap в PIL Image
                    mode = "RGBA" if pix.alpha else "RGB"
                    img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                    
                    try:
                        ocr_text = pytesseract.image_to_string(img, lang="rus+eng")
                        if ocr_text:
                            page_text += ocr_text + "\n"
                    except Exception as ocr_e:
                        logger.warning(f"OCR failed on page {i} of {file_path}: {ocr_e}")

                # 4. Добавляем номер страницы к каждому абзацу (включая таблицы и OCR-текст)
                if page_text.strip():
                    paragraphs = page_text.split('\n\n')
                    for para in paragraphs:
                        if para.strip():
                            text += f"[стр. {i}] " + para.strip() + "\n\n"
                            
            if doc_fitz:
                doc_fitz.close()
                
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path}: {e}")
        
    return text

def parse_docx(file_path: str) -> str:
    text = ""
    try:
        doc = DocxDocument(file_path)
        
        # Читаем абзацы
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text.strip() + "\n\n"
                
        # Читаем таблицы
        for table in doc.tables:
            md_table = ""
            for row_idx, row in enumerate(table.rows):
                row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                md_table += "| " + " | ".join(row_data) + " |\n"
                if row_idx == 0:
                    md_table += "|" + "|".join(["---"] * len(row.cells)) + "|\n"
            text += "\n" + md_table + "\n\n"
            
    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}")
    return text

def parse_excel(file_path: str) -> str:
    text = ""
    try:
        xls = pd.ExcelFile(file_path)
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            df.dropna(how='all', inplace=True)
            df.dropna(axis=1, how='all', inplace=True)
            
            if not df.empty:
                text += f"Лист: {sheet_name}\n"
                text += df.to_markdown(index=False) + "\n\n"
    except Exception as e:
        logger.error(f"Error parsing Excel {file_path}: {e}")
    return text

def parse_csv(file_path: str) -> str:
    text = ""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                row_str = " | ".join(row)
                if row_str:
                    text += row_str + "\n"
    except Exception as e:
        logger.error(f"Error parsing CSV {file_path}: {e}")
    return text

def parse_image(file_path: str) -> str:
    text = ""
    try:
        img = Image.open(file_path)
        logger.info(f"Running OCR on image {os.path.basename(file_path)}...")
        ocr_text = pytesseract.image_to_string(img, lang="rus+eng")
        if ocr_text.strip():
            text += ocr_text.strip() + "\n\n"
    except Exception as e:
        logger.error(f"Error parsing Image {file_path}: {e}")
    return text

def process_document(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    content = ""
    
    if ext == ".pdf":
        content = parse_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        content = parse_docx(file_path)
    elif ext in [".xlsx", ".xls"]:
        content = parse_excel(file_path)
    elif ext == ".csv":
        content = parse_csv(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
        content = parse_image(file_path)
    else:
        logger.warning(f"Unsupported extension: {ext}")
        
    return {
        "filename": os.path.basename(file_path),
        "content": content.strip()
    }
