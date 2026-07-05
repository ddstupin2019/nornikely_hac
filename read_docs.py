import openpyxl
import docx

def read_docx(path):
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip() != ""])

def read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True, max_row=10):
        rows.append("\t".join([str(c) if c is not None else "" for c in row]))
    return "\n".join(rows)

if __name__ == '__main__':
    print("--- DOCX ---")
    try:
        print(read_docx(r"c:\Users\ddstu\Desktop\nor_hac\test\exepmle\Как читать отчет института по хвостам.docx")[:1000])
    except Exception as e:
        print(f"Error reading docx: {e}")

    print("\n--- XLSX 1 ---")
    try:
        print(read_xlsx(r"c:\Users\ddstu\Desktop\nor_hac\test\exepmle\Пример 1\Хвосты КГМК.xlsx"))
    except Exception as e:
        print(f"Error reading xlsx 1: {e}")

    print("\n--- XLSX 4 ---")
    try:
        print(read_xlsx(r"c:\Users\ddstu\Desktop\nor_hac\test\exepmle\Пример 4\Хвосты ТОФ_2.xlsx"))
    except Exception as e:
        print(f"Error reading xlsx 4: {e}")
